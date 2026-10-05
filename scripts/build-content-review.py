#!/usr/bin/env python3
"""Builds the expert-review workbook from the app's real content.

  cd packages/shared && npx tsx scripts/dump-content.ts > /tmp/content.json
  python3 scripts/build-content-review.py /tmp/content.json docs/naqa-content-review.xlsx
"""
import json, sys
from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

src, out = sys.argv[1], sys.argv[2]
data = json.load(open(src, encoding='utf-8'))

STATUSES = ['صحيحة', 'تحتاج تعديل', 'خطأ', 'غير متأكد']
NAVY = '1B2A8F'; SOFT = 'E4E8FD'
thin = Side(style='thin', color='C9D0E6')
border = Border(left=thin, right=thin, top=thin, bottom=thin)

wb = Workbook()
ws0 = wb.active
ws0.title = 'اقرأ أولًا'
ws0.sheet_view.rightToLeft = True

# Which sheets need which kind of reviewer
who = {
  'الملابس الخليجية': 'عامل مغسلة أو خيّاط يعرف الملابس الخليجية',
  'إزالة البقع': 'مختص تنظيف ملابس / مغسلة',
  'رموز الغسيل': 'مختص أقمشة (مرجع: معيار رموز العناية الدولي)',
  'الغسالات والمجففات': 'فني أجهزة منزلية',
  'القطع': 'مختص أقمشة أو مغسلة',
  'الأقمشة': 'مهندس نسيج أو مختص أقمشة',
  'الألوان': 'مختص تنظيف ملابس',
  'قواعد الغسيل الأساسية': 'مختص أقمشة — الأهم لسلامة الملابس',
  'المنظفات والمنعّمات': 'مختص تنظيف ملابس أو كيميائي منظفات',
  'المنتجات': 'مختص منتجات التنظيف',
  'رسائل وتنبيهات': 'أي مراجع',
  'النصوص القانونية': 'محامٍ (وليس مختص أقمشة)',
}

lines = [
  ('مراجعة محتوى تطبيق نقاء', 'title'),
  ('', ''),
  ('الهدف: التأكد من أن كل نصيحة غسيل في التطبيق صحيحة وآمنة قبل أن تصل إلى الناس. كُتب هذا المحتوى من معلومات عامة ولم يُجرَّب ولم يراجعه مختص بعد.', ''),
  ('', ''),
  ('طريقة المراجعة', 'h'),
  ('1) افتح كل ورقة عمل (في أسفل الشاشة) وعليك قراءة عمود «النص بالعربية» سطرًا سطرًا.', ''),
  ('2) في عمود «التقييم» اختر من القائمة: صحيحة / تحتاج تعديل / خطأ / غير متأكد.', ''),
  ('3) إن اخترت «تحتاج تعديل» أو «خطأ» فاكتب الصيغة الصحيحة في عمود «التصحيح المقترح» (يكفي بالعربية).', ''),
  ('4) اكتب في «المصدر / المرجع» من أين تعرف ذلك إن أمكن (خبرة عمل، ملصق شركة، معيار…).', ''),
  ('5) كل ما كان فيه خطر على القطعة أو على الصحة (مبيّضات، كحول، ماء الأكسجين، خلط مواد) نريد رأيك فيه بالذات.', ''),
  ('', ''),
  ('معنى التقييم', 'h'),
  ('صحيحة: النص دقيق وآمن كما هو.   تحتاج تعديل: المعنى سليم لكن الصياغة أو التفصيل (وقت، حرارة، جرعة) غير دقيق.', ''),
  ('خطأ: النصيحة غير صحيحة أو قد تضر.   غير متأكد: يحتاج مراجعة مختص آخر.', ''),
  ('', ''),
  ('ملاحظات', 'h'),
  ('• ورقة «قواعد الغسيل الأساسية» أهم الأوراق: هي التي يُبنى عليها برنامج الغسيل والحرارة لكل قماش.', ''),
  ('• ورقة «النصوص القانونية» مسودة تحتاج محاميًا، وليست من اختصاص مراجع الأقمشة.', ''),
  ('• عمود «الرقم المرجعي» يربط كل سطر بمكانه في التطبيق، فلا تغيّره.', ''),
  ('', ''),
  ('المراجع', 'h'),
  ('الاسم:', ''), ('التخصص / الجهة:', ''), ('التاريخ:', ''),
]
for i, (text, kind) in enumerate(lines, start=1):
    c = ws0.cell(row=i, column=1, value=text)
    c.alignment = Alignment(wrap_text=True, vertical='top', horizontal='right')
    if kind == 'title': c.font = Font(bold=True, size=18, color=NAVY)
    elif kind == 'h': c.font = Font(bold=True, size=13, color=NAVY)
    else: c.font = Font(size=12)
ws0.column_dimensions['A'].width = 110
for i in range(len(lines) - 2, len(lines) + 1):
    ws0.cell(row=i, column=2).border = border
ws0.column_dimensions['B'].width = 40

summary_rows = []
headers = ['الرقم المرجعي', 'القسم', 'العنصر', 'النص بالعربية', 'النص بالإنجليزية', 'التقييم', 'التصحيح المقترح', 'المصدر / المرجع', 'ملاحظات المراجع']
widths = [22, 26, 24, 62, 48, 14, 48, 28, 32]

for name, rows in data.items():
    ws = wb.create_sheet(title=name[:31])
    ws.sheet_view.rightToLeft = True
    ws.sheet_properties.tabColor = NAVY
    ws.cell(row=1, column=1, value=f'المراجع المناسب: {who.get(name, "مختص")}')
    ws.cell(row=1, column=1).font = Font(bold=True, color=NAVY, size=12)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    for j, h in enumerate(headers, start=1):
        c = ws.cell(row=2, column=j, value=h)
        c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor=NAVY)
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True); c.border = border
        ws.column_dimensions[get_column_letter(j)].width = widths[j - 1]
    for i, r in enumerate(rows, start=3):
        vals = [r['ref'], r['section'], r['item'], r['ar'], r['en'], None, None, None, None]
        for j, v in enumerate(vals, start=1):
            c = ws.cell(row=i, column=j, value=v)
            c.alignment = Alignment(wrap_text=True, vertical='top', horizontal='left' if j == 5 else 'right')
            c.border = border
            if j in (6, 7, 8, 9): c.fill = PatternFill('solid', fgColor=SOFT)
    last = len(rows) + 2
    dv = DataValidation(type='list', formula1='"' + ','.join(STATUSES) + '"', allow_blank=True, showErrorMessage=True,
                        errorTitle='اختر من القائمة', error='اختر: ' + ' / '.join(STATUSES))
    ws.add_data_validation(dv); dv.add(f'F3:F{last}')
    for status, color in [('صحيحة', 'C6EFCE'), ('تحتاج تعديل', 'FFEB9C'), ('خطأ', 'FFC7CE'), ('غير متأكد', 'D9D9D9')]:
        ws.conditional_formatting.add(f'F3:F{last}', FormulaRule(formula=[f'$F3="{status}"'], fill=PatternFill('solid', bgColor=color, fgColor=color)))
    ws.freeze_panes = 'D3'
    ws.auto_filter.ref = f'A2:{get_column_letter(len(headers))}{last}'
    summary_rows.append((name, ws.title, last))

# Progress summary
ws = wb.create_sheet(title='ملخص التقدم', index=1)
ws.sheet_view.rightToLeft = True
for j, h in enumerate(['الورقة', 'عدد البنود', 'صحيحة', 'تحتاج تعديل', 'خطأ', 'غير متأكد', 'لم تُراجع'], start=1):
    c = ws.cell(row=1, column=j, value=h); c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor=NAVY)
    c.alignment = Alignment(horizontal='center'); c.border = border
    ws.column_dimensions[get_column_letter(j)].width = 22 if j == 1 else 14
for i, (name, title, last) in enumerate(summary_rows, start=2):
    rng = f"'{title}'!F3:F{last}"
    ws.cell(row=i, column=1, value=name)
    ws.cell(row=i, column=2, value=f'=ROWS({rng})')
    for j, st in enumerate(STATUSES, start=3):
        ws.cell(row=i, column=j, value=f'=COUNTIF({rng},"{st}")')
    ws.cell(row=i, column=7, value=f'=B{i}-SUM(C{i}:F{i})')
    for j in range(1, 8): ws.cell(row=i, column=j).border = border
t = len(summary_rows) + 2
ws.cell(row=t, column=1, value='المجموع').font = Font(bold=True)
for j in range(2, 8):
    col = get_column_letter(j)
    ws.cell(row=t, column=j, value=f'=SUM({col}2:{col}{t-1})').font = Font(bold=True)
    ws.cell(row=t, column=j).border = border
ws.cell(row=t, column=1).border = border

wb.save(out)
print('saved', out, sum(len(v) for v in data.values()), 'rows')
